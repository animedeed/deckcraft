"""Сборка .pptx из списка слайдов: фигуры и текст через python-pptx, затем постобработка XML —
переходы Morph, шкала анимаций (появление anim_* и петли loop_*), цвета и шрифты темы.
"""

from __future__ import annotations

import io
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Inches, Pt

from .anim import Effect, FxCall, motion_path, timing_xml, transition_xml
from .presets import Fx, duration_ms
from .model import Loop, O, Slide, apply_text_policy, parse_line, resolve
from .profile import DEFAULT, DeckProfile
from .theme import SLIDE_H, SLIDE_W, Theme

NS = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
A = "{%s}" % NS["a"]
SHAPES = {"fold": MSO_SHAPE.FOLDED_CORNER, "rect": MSO_SHAPE.RECTANGLE, "rrect": MSO_SHAPE.ROUNDED_RECTANGLE, "oval": MSO_SHAPE.OVAL,
          "arc": MSO_SHAPE.BLOCK_ARC, "chevron": MSO_SHAPE.CHEVRON, "tri": MSO_SHAPE.ISOSCELES_TRIANGLE,
          "pent": MSO_SHAPE.PENTAGON, "hex": MSO_SHAPE.HEXAGON, "donut": MSO_SHAPE.DONUT}
DASH = {"dash": MSO_LINE_DASH_STYLE.DASH, "dot": MSO_LINE_DASH_STYLE.ROUND_DOT,
        "sysDash": MSO_LINE_DASH_STYLE.SQUARE_DOT, "lgDash": MSO_LINE_DASH_STYLE.LONG_DASH}
APPEAR_DUR = 500
APPEAR_STEP = 380
LOOP_START = 150  # мс после перехода


# ----------------------------------------------------------------------------- петли -> объекты и эффекты

def expand_loops(slide: Slide, objs: list[O], theme: Theme) -> tuple[list[O], list[tuple[str, Loop, dict]]]:
    """Добавляет объекты-носители петель (точки потоков, кольца) и возвращает их будущие эффекты."""
    by_name = {o.name: o for o in objs}
    extra, plans = [], []
    for k, lp in enumerate(slide.loops):
        base = f"loop_{lp.name}_{k}"
        if lp.kind == "flow":
            pts = list(reversed(lp.points)) if lp.reverse else list(lp.points)
            for i in range(lp.n):
                name = f"{base}_dot{i}"
                extra.append(O(name, "oval", pts[0][0] - lp.size / 2, pts[0][1] - lp.size / 2, lp.size, lp.size,
                               fill=lp.color, z=lp.z, meta={"loop": True}))
                plans.append((name, lp, {"phase": i / lp.n, "points": pts}))
        elif lp.kind == "ripple":
            t = by_name[lp.target]
            name = f"{base}_ring"
            extra.append(O(name, "oval", t.x, t.y, t.w, t.h, fill=None, ln=lp.color, lw=2.0,
                           z=t.z - 0.5, meta={"loop": True}))
            plans.append((name, lp, {}))
        else:
            if lp.target not in by_name:
                raise KeyError(f"петля {lp.kind}: нет объекта {lp.target} на слайде «{slide.title}»")
            plans.append((lp.target, lp, {}))
    return extra, plans


def schedule_fx(fxs: list[Fx], spids: dict[str, int], has_text: set[str], theme: Theme,
                title: str = "") -> list[FxCall]:
    """Старт как в панели анимации: with — вместе с предыдущим, after — после всех предыдущих (кроме
    повторяющихся), at — от начала шкалы. Цвет в options может быть ролью темы."""
    calls, prev_start, end = [], 0, 0
    for f in fxs:
        if f.target not in spids:
            raise KeyError(f"Fx {f.key}: нет объекта {f.target!r} на слайде «{title}»")
        if "color" in f.options:
            f = Fx(**{**f.__dict__, "options": {**f.options, "color": theme.c(f.options["color"])}})
        lag = int(round(f.delay * 1000))
        start = {"with": prev_start + lag, "after": end + lag, "at": lag}[f.start]
        calls.append(FxCall(f, spids[f.target], start, f.target in has_text))
        if not f.repeat:
            end = max(end, start + duration_ms(f))
        prev_start = start
    return calls


def loop_effects(plans, spids: dict[str, int], base_delay: int) -> list[Effect]:
    effects = []
    for name, lp, info in plans:
        period = int(lp.period * 1000)
        delay = base_delay + int(lp.start * 1000)
        if lp.kind == "flow":
            d = delay + int(info["phase"] * period)
            effects.append(Effect("show", spids[name], d, 1))  # точка видна только с момента своего старта
            effects.append(Effect("flow", spids[name], d, period, repeat=True,
                                  path=motion_path(info["points"], SLIDE_W, SLIDE_H)))
        elif lp.kind == "pulse":
            effects.append(Effect("pulse", spids[name], delay, period // 2, repeat=True, scale=lp.scale))
        elif lp.kind == "ripple":
            effects.append(Effect("show", spids[name], delay, 1))
            effects.append(Effect("ripple", spids[name], delay, period, repeat=True, scale=lp.scale))
        elif lp.kind == "spin":  # reverse=True — против часовой (встречные шестерни)
            effects.append(Effect("spin", spids[name], delay, period, repeat=True, scale=-1.0 if lp.reverse else 1.0))
        elif lp.kind == "blink":
            effects.append(Effect(lp.kind, spids[name], delay, period, repeat=True))
        else:
            raise ValueError(lp.kind)
    return effects


# ----------------------------------------------------------------------------- текст

def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def txbody_xml(o: O, theme: Theme, inset: float = 0.0) -> str:
    anchor = {"t": "t", "m": "ctr", "b": "b"}[o.va]
    algn = {"l": "l", "c": "ctr", "r": "r"}[o.al]
    ins = int(Inches(inset))
    paras = []
    prepared = o.meta.get("runs")  # список строк, каждая — список Run (панель кода и т.п.)
    for li, line in enumerate(o.lines):
        bullet = o.bullets and line.strip() != ""
        ppr = f'<a:pPr algn="{algn}"' + (' marL="228600" indent="-228600"' if bullet else "") + ">"
        if o.meta.get("line_pts"):  # точный шаг строк (одометр, колонки цифр)
            ppr += f'<a:lnSpc><a:spcPts val="{int(o.meta["line_pts"] * 100)}"/></a:lnSpc>'
        else:
            ppr += f'<a:lnSpc><a:spcPct val="{int(o.spacing * 100000)}"/></a:lnSpc>'
        if o.para_gap:
            ppr += f'<a:spcAft><a:spcPts val="{int(o.para_gap * 100)}"/></a:spcAft>'
        ppr += ('<a:buFont typeface="Arial"/><a:buChar char="•"/>' if bullet else "<a:buNone/>") + "</a:pPr>"
        runs = []
        for run in (prepared[li] if prepared else parse_line(line)):
            if run.text == "":
                continue
            color = theme.c(run.color or o.col)
            face = theme.font("mono" if run.mono else o.font)
            bold = o.b if run.bold is None else run.bold
            attrs = f'lang="ru-RU" sz="{int(round(o.sz * 100))}" b="{1 if bold else 0}" dirty="0"'
            if o.caps:
                attrs += ' cap="all"'
            if o.tracking:
                attrs += f' spc="{int(o.tracking * 100)}"'
            if run.strike:
                attrs += ' strike="sngStrike"'
            runs.append(f'<a:r><a:rPr {attrs}><a:solidFill><a:srgbClr val="{color}"/></a:solidFill>'
                        f'<a:latin typeface="{face}"/><a:cs typeface="{face}"/></a:rPr><a:t>{escape(run.text)}</a:t></a:r>')
        end = f'<a:endParaRPr lang="ru-RU" sz="{int(round(o.sz * 100))}" dirty="0"/>'
        paras.append(f"<a:p>{ppr}{''.join(runs)}{end}</a:p>")
    return (f'<p:txBody xmlns:p="{NS["p"]}" xmlns:a="{NS["a"]}"><a:bodyPr wrap="square" lIns="{ins}" tIns="{ins}" '
            f'rIns="{ins}" bIns="{ins}" anchor="{anchor}" rtlCol="0"><a:noAutofit/></a:bodyPr><a:lstStyle/>'
            f'{"".join(paras) or "<a:p/>"}</p:txBody>')


def _set_txbody(shape, xml: str) -> None:
    el = shape._element
    old = el.find("{%s}txBody" % NS["p"])
    new = etree.fromstring(xml)
    if old is not None:
        el.replace(old, new)
    else:
        el.append(new)


def _alpha(parent_fill, alpha: float) -> None:
    if alpha >= 0.999:
        return
    clr = parent_fill.find(A + "srgbClr")
    if clr is not None:
        etree.SubElement(clr, A + "alpha").set("val", str(int(alpha * 100000)))


def _style_fill(shape, o: O, theme: Theme) -> None:
    if o.fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(theme.c(o.fill))
        sp_pr = shape._element.spPr
        _alpha(sp_pr.find(A + "solidFill"), o.fa)


def _style_line(shape, o: O, theme: Theme) -> None:
    if o.ln is None or o.lw <= 0:
        shape.line.fill.background()
        return
    shape.line.color.rgb = _rgb(theme.c(o.ln))
    shape.line.width = Pt(o.lw)
    if o.dash:
        shape.line.dash_style = DASH[o.dash]
    ln = shape._element.spPr.find(A + "ln")
    _alpha(ln.find(A + "solidFill"), o.la)
    if o.head:
        tail = etree.SubElement(ln, A + "tailEnd")
        tail.set("type", o.head)
        tail.set("w", "med")
        tail.set("len", "med")


def _set_adj(shape, values: dict[str, int]) -> None:
    geom = shape._element.spPr.find(A + "prstGeom")
    av = geom.find(A + "avLst")
    if av is None:
        av = etree.SubElement(geom, A + "avLst")
    for child in list(av):
        av.remove(child)
    for name, val in values.items():
        gd = etree.SubElement(av, A + "gd")
        gd.set("name", name)
        gd.set("fmla", f"val {int(val)}")


def _grad_fill_xml(grad: dict, theme: Theme) -> str:
    gs = []
    for st in grad["stops"]:
        pos, color = st[0], st[1]
        alpha = st[2] if len(st) > 2 else 1.0
        a = f'<a:alpha val="{int(alpha * 100000)}"/>' if alpha < 0.999 else ""
        gs.append(f'<a:gs pos="{int(pos * 100000)}"><a:srgbClr val="{theme.c(color)}">{a}</a:srgbClr></a:gs>')
    if grad.get("kind", "linear") == "linear":
        shade = f'<a:lin ang="{int(grad.get("angle", 90) * 60000) % 21600000}" scaled="0"/>'
    else:
        shade = ('<a:path path="circle"><a:fillToRect l="50000" t="50000" r="50000" b="50000"/></a:path>'
                 if grad["kind"] == "radial" else
                 '<a:path path="rect"><a:fillToRect l="50000" t="50000" r="50000" b="50000"/></a:path>')
    return f'<a:gradFill xmlns:a="{NS["a"]}" rotWithShape="1"><a:gsLst>{"".join(gs)}</a:gsLst>{shade}</a:gradFill>'


def _apply_grad(sh, o: O, theme: Theme) -> None:
    sp_pr = sh._element.spPr
    old = None
    for tag in ("solidFill", "noFill", "gradFill"):
        old = sp_pr.find(A + tag)
        if old is not None:
            break
    new = etree.fromstring(_grad_fill_xml(o.meta["grad"], theme))
    if old is not None:
        sp_pr.replace(old, new)
    else:
        geom = sp_pr.find(A + "custGeom")
        geom = geom if geom is not None else sp_pr.find(A + "prstGeom")
        geom.addnext(new)


def _apply_3d(sh, o: O, theme: Theme) -> None:
    d = o.meta["d3"]
    deg = lambda v: int(round((v % 360) * 60000))  # noqa: E731
    fov = f' fov="{int(d.get("fov", 45) * 60000)}"' if d.get("camera", "").startswith("perspective") else ""
    scene = (f'<a:scene3d xmlns:a="{NS["a"]}"><a:camera prst="{d.get("camera", "perspectiveFront")}"{fov}>'
             f'<a:rot lat="{deg(d.get("y", 0))}" lon="{deg(d.get("x", 0))}" rev="{deg(d.get("z", 0))}"/></a:camera>'
             f'<a:lightRig rig="{d.get("light", "threePt")}" dir="t"/></a:scene3d>')
    attrs = ""
    if d.get("depth"):
        attrs += f' extrusionH="{int(Pt(d["depth"]))}"'
    if d.get("material"):
        attrs += f' prstMaterial="{d["material"]}"'
    inner = ""
    if d.get("bevel"):
        b = int(Pt(d["bevel"]))
        inner += f'<a:bevelT w="{b}" h="{b}"/>'
    if d.get("depth"):
        inner += f'<a:extrusionClr><a:srgbClr val="{theme.c(d.get("depth_color") or o.fill or "ink2")}"/></a:extrusionClr>'
    sp3d = f'<a:sp3d xmlns:a="{NS["a"]}"{attrs}>{inner}</a:sp3d>'
    sp_pr = sh._element.spPr
    ext = sp_pr.find(A + "extLst")
    for xml in (scene, sp3d):
        el = etree.fromstring(xml)
        if ext is not None:
            ext.addprevious(el)
        else:
            sp_pr.append(el)


def _set_custgeom(sh, o: O) -> None:
    from .vector import custgeom_xml
    sp_pr = sh._element.spPr
    old = sp_pr.find(A + "prstGeom")
    sp_pr.replace(old, etree.fromstring(custgeom_xml(o.meta["vec"])))


def add_object(slide_obj, o: O, theme: Theme):
    shapes = slide_obj.shapes
    if o.kind in ("line", "arrow"):
        x2, y2 = o.x + o.w, o.y + o.h
        sh = shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(o.x), Inches(o.y), Inches(x2), Inches(y2))
        oo = o if o.kind == "line" or o.head else o.copy(head="triangle")
        _style_line(sh, oo.copy(lw=oo.lw or 1.5, ln=oo.ln or "ink2"), theme)
    elif o.kind == "img":
        sh = shapes.add_picture(o.img, Inches(o.x), Inches(o.y), Inches(o.w), Inches(o.h))
        if o.crop:
            sh.crop_left, sh.crop_top, sh.crop_right, sh.crop_bottom = o.crop
    elif o.kind == "text":
        sh = shapes.add_textbox(Inches(o.x), Inches(o.y), Inches(o.w), Inches(o.h))
        if o.fill is not None:
            _style_fill(sh, o, theme)
        _set_txbody(sh, txbody_xml(o, theme, o.meta.get("inset", 0.0)))
    else:
        kind = "rect" if o.kind == "vec" else o.kind
        sh = shapes.add_shape(SHAPES[kind], Inches(o.x), Inches(o.y), Inches(max(o.w, 0.001)), Inches(max(o.h, 0.001)))
        if o.kind == "vec":
            _set_custgeom(sh, o)
        _style_fill(sh, o, theme)
        _style_line(sh, o, theme)
        if o.kind == "rrect":
            r = theme.radius if o.radius is None else o.radius
            _set_adj(sh, {"adj": min(50000, int(r / max(min(o.w, o.h), 1e-6) * 100000))})
        if o.adj:
            _set_adj(sh, o.adj)
        if o.txt is not None:
            _set_txbody(sh, txbody_xml(o, theme, o.meta.get("inset", 0.12)))
    if o.meta.get("grad") and o.kind not in ("line", "arrow", "img"):
        _apply_grad(sh, o, theme)
    if o.meta.get("d3") and o.kind not in ("line", "arrow"):
        _apply_3d(sh, o, theme)
    style = sh._element.find("{%s}style" % NS["p"])
    if style is not None:  # встроенный стиль фигуры добавляет тень и цвета темы — всё задаём явно
        sh._element.remove(style)
    sh.name = o.name
    if o.rot:
        sh.rotation = o.rot
    return sh


# ----------------------------------------------------------------------------- тема PowerPoint

def _patch_theme(xml: bytes, theme: Theme) -> bytes:
    root = etree.fromstring(xml)
    scheme = root.find(".//a:clrScheme", NS)
    mapping = {"dk1": "ink", "lt1": "bg", "dk2": "ink2", "lt2": "surface", "accent1": "accent", "accent2": "accent2",
               "accent3": "accent3", "accent4": "good", "accent5": "warn", "accent6": "bad", "hlink": "accent",
               "folHlink": "muted"}
    if scheme is not None:
        scheme.set("name", f"morphkit-{theme.name}")
        for tag, role in mapping.items():
            node = scheme.find(f"a:{tag}", NS)
            if node is None or role not in theme.colors:
                continue
            for child in list(node):
                node.remove(child)
            etree.SubElement(node, A + "srgbClr").set("val", theme.colors[role])
    for font_tag, face in (("majorFont", theme.head_font), ("minorFont", theme.body_font)):
        latin = root.find(f".//a:{font_tag}/a:latin", NS)
        if latin is not None:
            latin.set("typeface", face)
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


# ----------------------------------------------------------------------------- сборка

def build_pptx(slides: list[Slide], theme: Theme, out: str | Path, title: str = "",
               profile: DeckProfile | None = None) -> dict:
    profile = profile or DEFAULT
    slides = apply_text_policy(slides, profile.text_policy)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)
    if title:
        prs.core_properties.title = title
    blank = prs.slide_layouts[6]
    plans_per_slide, resolved_per_slide = [], []
    for s in slides:
        objs = resolve(s)
        extra, plans = expand_loops(s, objs, theme)
        objs = sorted(objs + extra, key=lambda o: o.z)
        resolved_per_slide.append(objs)
        plans_per_slide.append(plans)
        sl = prs.slides.add_slide(blank)
        bg = sl.background.fill
        bg.solid()
        bg.fore_color.rgb = _rgb(theme.c(s.bg))
        for o in objs:
            add_object(sl, o, theme)
        if s.notes:
            sl.notes_slide.notes_text_frame.text = s.notes
    buf = io.BytesIO()
    prs.save(buf)

    zin = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    order = zin.namelist()
    items = {n: zin.read(n) for n in order}
    zin.close()
    report = []
    for i, s in enumerate(slides):
        name = f"ppt/slides/slide{i + 1}.xml"
        root = etree.fromstring(items[name])
        spids = {el.get("name"): int(el.get("id")) for el in root.iter("{%s}cNvPr" % NS["p"])}
        has_text = _text_names(root)
        base = 0  # шкала слайда стартует после окончания перехода (проверено записью показа в PowerPoint)
        effects: list = []
        anim_names = s.anims or [o.name for o in resolved_per_slide[i] if o.is_anim]
        for k, n in enumerate(anim_names):
            effects.append(Effect("appear", spids[n], base + 150 + k * APPEAR_STEP, APPEAR_DUR, has_text=n in has_text))
        effects += loop_effects(plans_per_slide[i], spids, base + LOOP_START)
        effects += schedule_fx(s.fx, spids, has_text, theme, s.title)
        anchor = root.find("p:clrMapOvr", NS)
        pos = list(root).index(anchor) + 1
        tr = transition_xml(s.transition if i > 0 else "none", s.dur, s.mode, s.auto, s.trans_opts)
        if tr:
            root.insert(pos, etree.fromstring(tr))
            pos += 1
        tm = timing_xml(effects)
        if tm:
            root.insert(pos, etree.fromstring(tm))
        items[name] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        report.append({"slide": i + 1, "title": s.title, "objects": len(resolved_per_slide[i]),
                       "appear": len(anim_names), "loops": len(s.loops), "fx": len(s.fx), "transition": s.transition if i else "—"})
    theme_name = next((n for n in order if n.startswith("ppt/theme/theme") and n.endswith(".xml")), None)
    if theme_name:
        items[theme_name] = _patch_theme(items[theme_name], theme)
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in order:
            zout.writestr(n, items[n])
    return {"path": str(out), "slides": report, "resolved": resolved_per_slide}


def _text_names(root) -> set[str]:
    names = set()
    for sp in root.iter("{%s}sp" % NS["p"]):
        body = sp.find("{%s}txBody" % NS["p"])
        if body is not None and "".join(body.itertext()).strip():
            names.add(sp.find("{%s}nvSpPr/{%s}cNvPr" % (NS["p"], NS["p"])).get("name"))
    return names
